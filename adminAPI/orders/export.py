"""Order Excel export helpers."""

import logging
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from orders.models import Order

logger = logging.getLogger(__name__)

EXPORT_HEADERS = [
    '订单号',
    '下单时间',
    '客户名',
    '商品明细',
    '总金额',
    '支付方式',
    '收货地址',
    '订单状态',
    '物流单号',
]

STATUS_LABELS = dict(Order.STATUS_CHOICES)


def _format_datetime(value: datetime | None) -> str:
    """Format datetime as YYYY-MM-DD HH:mm:ss."""
    if not value:
        return ''
    if timezone.is_naive(value):
        value = timezone.make_aware(value, timezone.get_current_timezone())
    local_value = timezone.localtime(value)
    return local_value.strftime('%Y-%m-%d %H:%M:%S')


def _format_amount(value: Decimal | float | int | None) -> float:
    """Format amount with two decimal places."""
    if value is None:
        return 0.00
    return float(Decimal(value).quantize(Decimal('0.01')))


def _format_payment_method(order: Order) -> str:
    """Derive payment method label from order status."""
    if order.status == Order.STATUS_PENDING:
        return '未支付'
    return '在线支付'


def _format_items(order: Order) -> str:
    """Format order line items as readable text."""
    parts = []
    for item in order.items.all():
        parts.append(f'{item.product.name} x{item.quantity}')
    return '；'.join(parts)


def _auto_fit_columns(worksheet) -> None:
    """Adjust column widths based on content."""
    for column_cells in worksheet.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            cell_value = '' if cell.value is None else str(cell.value)
            max_length = max(max_length, len(cell_value))
        worksheet.column_dimensions[get_column_letter(column)].width = min(max_length + 2, 50)


def build_orders_workbook(orders) -> BytesIO:
    """Build Excel workbook for order export."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '订单列表'

    header_font = Font(bold=True)
    for col_index, header in enumerate(EXPORT_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.font = header_font

    for row_index, order in enumerate(orders, start=2):
        worksheet.cell(row=row_index, column=1, value=order.order_no)
        worksheet.cell(row=row_index, column=2, value=_format_datetime(order.created_at))
        worksheet.cell(row=row_index, column=3, value=order.customer.name)
        worksheet.cell(row=row_index, column=4, value=_format_items(order))
        worksheet.cell(row=row_index, column=5, value=_format_amount(order.total_amount))
        worksheet.cell(row=row_index, column=6, value=_format_payment_method(order))
        worksheet.cell(row=row_index, column=7, value=order.address)
        worksheet.cell(row=row_index, column=8, value=STATUS_LABELS.get(order.status, order.status))
        worksheet.cell(row=row_index, column=9, value=order.logistics_no or '')

    _auto_fit_columns(worksheet)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
